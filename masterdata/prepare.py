"""Build pre-match predictors for every workbook row and write the filled workbook.

Reads raw_history.csv / crosscheck.csv / friendly_lists.csv produced by collect.py.
Writes data_wrangling_filled.xlsx (original workbook is only read) and COMPLETENESS.md.

Source-match rule: every row's predictors come from exactly ONE previous match played by that team
(user, 2026-10-15). No historical averaging. Friendly-only approach (branch data-friendly-approach, user
decision 2026-10-16 after a team meeting: one match type for all openers):
  * opener row  -> the team's last international friendly before the opener; if that friendly has no
                   published team statistics on any source, the team's most recent friendly that has them
  * later rows  -> the team's previous WC 2026 match
  * if that one match lacks a stat, the cell stays blank; an older match is never used instead
Counts are per 90 nominal minutes (x 90/120 when that match went to extra time); percentages are the
match's own ratio. Saving when the opponent had 0 shots on target is coded 0 (user rule, 2026-10-16), and
saving is capped at 100 when a source reports more saves than opponent shots on target (data check, 2026-10-17).

Headers (copy only; the original workbook keeps its names): counts carry a _per90 suffix (see RENAMED), and
a new last column 'previous_match_duration' gives the source match's length: 90, or 120 after extra time.

Exception: 'FIFA ranking points' (replaced 'km running' on 2026-10-16 at the user's request) is not a match
statistic. It is the team's total points in the last official FIFA ranking published before kickoff
(11 Jun 2026 for every row; FIFA published no ranking during the tournament). Read from fifa_ranking.csv.
"""
import datetime as dt
import math
import sys

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import (FIFA_CALENDAR, FIFA_RANKING_PAGE, INTERMEDIATE, OUTPUT, RANKING_CSV, RAW_CSV, WB_TO_FIFA,
                    WORKBOOK_IN, WORKBOOK_OUT, fetch_json)

N_WINDOW = 1                             # exactly one previous match per row (no averaging)
STATS_PRESENCE = ["shots_total", "possession_pct", "fouls_committed", "passes_att"]   # any present = has stats
WT_WC = "previous WC match"
WT_FR = "opener: last friendly"
WT_FR_EARLIER = "opener: latest friendly with statistics (last friendly has none)"
FK_MAX_MAE, FK_MIN_R = 1.0, 0.9          # acceptance rule for the derived free-kick measure

# workbook column (header in this copy) -> (internal key, unit label)
COLUMNS = {
    "free_kicks_won_per90": ("free_kick", "free kicks won per 90"),
    "Shot accurrancy": ("shot_accuracy", "% of shots on target"),
    "Pass acurracy": ("pass_accuracy", "% of passes completed"),
    "possession": ("possession", "% ball possession"),
    "saving": ("saving", "% of opponent shots on target saved (0 when the opponent had no shot on target)"),
    "fouls_committed_per90": ("foul", "fouls committed per 90"),
    "tackles_attempted_per90": ("tackle", "tackles (attempted) per 90"),
    "corners_per90": ("corner", "corners won per 90"),
    "FIFA ranking points": ("fifa_ranking_points", "FIFA men's world ranking total points, last release before "
                                                   "kickoff"),
    "goals_conceded_per90": ("goal_conceded", "goals conceded per 90 (previous match, not this one)"),
    "cards_per90": ("card", "yellow + red cards per 90"),
    "shots_per90": ("shot", "total shots per 90"),
}
# header in this copy -> header in the original workbook (user request 2026-10-16; column positions unchanged)
RENAMED = {"free_kicks_won_per90": "free kick", "fouls_committed_per90": "foul",
           "tackles_attempted_per90": "tackle", "corners_per90": "corner", "FIFA ranking points": "km running",
           "goals_conceded_per90": "goal conced", "cards_per90": "card", "shots_per90": "shot"}
DURATION = "previous_match_duration"     # new last column in this copy: 90, or 120 if the source match had ET
RANKING = "fifa_ranking_points"                   # not from the source match (see module docstring)


# raw column(s) that each predictor needs per match (used to name matches lacking data)
NEEDS = {"free_kick": ["opp_fouls_committed", "opp_offsides"], "shot_accuracy": ["shots_on_target", "shots_total"],
         "pass_accuracy": ["passes_cmp", "passes_att"], "possession": ["possession_pct"],
         "saving": ["saves", "opp_shots_on_target"], "foul": ["fouls_committed"], "tackle": ["tackles"],
         "corner": ["corners"], "goal_conceded": ["goals_against"],
         "card": ["yellow_cards", "red_cards"], "shot": ["shots_total"]}


def lacking(win, key):
    """Describe window matches without the needed raw stat."""
    bad = win[win[NEEDS[key]].isna().any(axis=1)]
    return "; ".join(f"{m['kickoff_utc'][:10]} vs {m['opponent']}" for _, m in bad.iterrows())


def per90(win, col):
    ok = win[win[col].notna()]
    if ok.empty:
        return None, 0
    return ok[col].sum() / ok["nominal_minutes"].sum() * 90, len(ok)


def pooled(win, num, den):
    ok = win[win[num].notna() & win[den].notna()]
    if ok.empty:
        return None, 0, "no usable matches"
    if ok[den].sum() == 0:
        if den == "opp_shots_on_target":
            # user rule (2026-10-16): no shot on target faced counts as 0, like a keeper who saved none
            return 0.0, len(ok), "0/0 coded as 0: opponent had 0 shots on target in that match"
        what = {"shots_total": "team had 0 shots", "passes_att": "team attempted 0 passes"}[den]
        return None, len(ok), f"undefined (0/0): {what} in that match"
    pct = 100 * ok[num].sum() / ok[den].sum()
    if den == "opp_shots_on_target" and pct > 100:
        # a keeper cannot save more shots than were on target; friendly saves come from FotMob, whose own
        # counts can disagree (data check 17 Oct 2026: Senegal 4 saves v Saudi Arabia's 3 on target, 0 conceded)
        return 100.0, len(ok), (f"capped at 100: source reports {int(ok[num].sum())} saves v "
                                f"{int(ok[den].sum())} opponent shots on target")
    return pct, len(ok), None


def aggregate(win, fk_ok):
    out, n, why = {}, {}, {}
    w = win.copy()
    w["fk_derived"] = w["opp_fouls_committed"] + w["opp_offsides"]
    w["cards"] = w["yellow_cards"] + w["red_cards"]
    fk_col = "fk_derived" if fk_ok else "fk_won_reported"
    for key, col in [("free_kick", fk_col), ("foul", "fouls_committed"), ("tackle", "tackles"),
                     ("corner", "corners"), ("goal_conceded", "goals_against"),
                     ("card", "cards"), ("shot", "shots_total")]:
        out[key], n[key] = per90(w, col)
    for key, num, den in [("shot_accuracy", "shots_on_target", "shots_total"),
                          ("pass_accuracy", "passes_cmp", "passes_att"),
                          ("saving", "saves", "opp_shots_on_target")]:
        out[key], n[key], why[key] = pooled(w, num, den)
    ok = w[w["possession_pct"].notna()]
    n["possession"] = len(ok)
    out["possession"] = (ok["possession_pct"] * ok["nominal_minutes"]).sum() / ok["nominal_minutes"].sum() \
        if len(ok) else None
    return out, n, why


def describe(m):
    ha = "H" if m["home_away"] == "home" else "A"
    tag = (f"WC#{int(m['wc_match_number'])}" if m["match_type"] == "WC2026" else
           "FRI" if m["match_type"] == "FRIENDLY2026" else "WCQ")
    et = " AET" if m["extra_time"] else ""
    return (f"{m['kickoff_utc'][:10]} {tag} vs {m['opponent']} ({ha}) {int(m['goals_for'])}-"
            f"{int(m['goals_against'])}{et} [FIFA {m['fifa_match_id']}]")


def parse_wb_date(v):
    """Workbook dates are day/month/year. Excel turned day<=12 into datetimes with day and month swapped."""
    if isinstance(v, dt.datetime):
        return dt.date(v.year, v.day, v.month), "Excel datetime with day/month swapped; corrected"
    d, m, y = (int(x) for x in str(v).strip().split("/"))
    return dt.date(y, m, d), "text dd/mm/yyyy; converted to date"


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--teams", default=None, help="comma-separated workbook team names (pilot)")
    args = ap.parse_args()
    raw = pd.read_csv(RAW_CSV, dtype={"fifa_match_id": str})
    raw["kickoff"] = pd.to_datetime(raw["kickoff_utc"], utc=True)
    cross = pd.read_csv(INTERMEDIATE / "crosscheck.csv")
    lists = pd.read_csv(INTERMEDIATE / "friendly_lists.csv")
    cal = {m["MatchNumber"]: m for m in fetch_json(FIFA_CALENDAR, "fifa_calendar")["Results"]}
    ranking = pd.read_csv(RANKING_CSV, dtype={"fifa_id_team": str})
    ranking["released"] = pd.to_datetime(ranking["release_utc"], utc=True)

    # ---- free-kick identity validation on WC matches (FIFA FreeKicks vs opp fouls + opp offsides)
    wc = raw[raw["match_type"] == "WC2026"].copy()
    wc["fk_derived"] = wc["opp_fouls_committed"] + wc["opp_offsides"]
    v = wc.dropna(subset=["fk_won_reported", "fk_derived"])
    fk_mae = float((v["fk_derived"] - v["fk_won_reported"]).abs().mean())
    fk_bias = float((v["fk_derived"] - v["fk_won_reported"]).mean())
    fk_r = float(np.corrcoef(v["fk_derived"], v["fk_won_reported"])[0, 1]) if len(v) > 2 else float("nan")
    fk_exact = float((v["fk_derived"] == v["fk_won_reported"]).mean())
    fk_ok = len(v) >= 20 and fk_mae <= FK_MAX_MAE and fk_r >= FK_MIN_R
    print(f"free-kick validation: n={len(v)} MAE={fk_mae:.2f} bias={fk_bias:+.2f} r={fk_r:.3f} "
          f"exact={fk_exact:.0%} -> {'derived accepted' if fk_ok else 'derived REJECTED'}")

    wbook = openpyxl.load_workbook(WORKBOOK_IN)
    ws = wbook["in"]
    header = {c.value: c.column for c in ws[1]}
    for new, old in RENAMED.items():               # rename in the copy only; column position unchanged
        header[new] = header.pop(old)
        ws.cell(1, header[new]).value = new
    header[DURATION] = ws.max_column + 1             # new last column (S), plain header like the others
    ws.cell(1, header[DURATION]).value = DURATION
    ws.column_dimensions[get_column_letter(header[DURATION])].width = len(DURATION) + 2
    all_wb_teams = {ws.cell(r, header["Team"]).value for r in range(2, ws.max_row + 1)}
    wanted_teams = ({t.strip() for t in args.teams.split(",")} if args.teams else all_wb_teams)

    out_rows, win_rows, miss_rows, date_log = [], [], [], []
    used_in = {}                       # (team, fifa_match_id) -> workbook rows that use this match
    for r in range(2, ws.max_row + 1):
        team_wb = ws.cell(r, header["Team"]).value
        num = ws.cell(r, header["MatchNumber"]).value
        fname = WB_TO_FIFA.get(team_wb, team_wb)
        m = cal[num]
        kickoff = pd.Timestamp(m["Date"])
        date_raw = ws.cell(r, header["Date"]).value
        date_fixed, how = parse_wb_date(date_raw)
        date_log.append({"excel_row": r, "match_number": num, "team": team_wb,
                         "round": ws.cell(r, header["Round"]).value, "original_value": str(date_raw),
                         "original_type": type(date_raw).__name__, "normalised_date": date_fixed.isoformat(),
                         "fifa_kickoff_utc": m["Date"], "how": how,
                         "matches_fifa_utc_date": date_fixed.isoformat() == m["Date"][:10]})
        is_home = m["Home"]["TeamName"][0]["Description"] == fname
        act = m["HomeTeamScore"] if is_home else m["AwayTeamScore"]
        id_team = m["Home"]["IdTeam"] if is_home else m["Away"]["IdTeam"]
        if team_wb not in wanted_teams:
            out_rows.append((r, None, None, act, date_fixed))
            continue
        team_raw = raw[raw["team"] == fname]
        prior_wc = team_raw[(team_raw["match_type"] == "WC2026") & (team_raw["kickoff"] < kickoff)] \
            .sort_values("kickoff", ascending=False)
        fr = team_raw[(team_raw["match_type"] == "FRIENDLY2026") & (team_raw["kickoff"] < kickoff)] \
            .sort_values("kickoff", ascending=False)
        opener = prior_wc.empty
        if not opener:
            win, wtype = prior_wc.head(N_WINDOW), WT_WC
        else:
            # last friendly; if it has no team statistics, the newest collected friendly that has them
            with_stats = fr[fr[STATS_PRESENCE].notna().any(axis=1)]
            if len(fr) and len(with_stats) and with_stats.index[0] != fr.index[0]:
                win, wtype = with_stats.head(N_WINDOW), WT_FR_EARLIER
            else:
                win, wtype = fr.head(N_WINDOW), WT_FR if len(fr) else "opener: no friendly found"
        assert (win["kickoff"] < kickoff).all(), (r, team_wb)
        assert (win["match_type"] == ("FRIENDLY2026" if opener else "WC2026")).all(), (r, team_wb)
        if win.empty:
            vals, ns, why = {k: None for k, _ in COLUMNS.values()}, {k: 0 for k, _ in COLUMNS.values()}, {}
        else:
            vals, ns, why = aggregate(win, fk_ok)
        # ranking: latest official release before kickoff, independent of the source match
        before = ranking[(ranking["released"] < kickoff) & (ranking["fifa_id_team"] == id_team)]
        rel = before.sort_values("released").iloc[-1] if len(before) else None
        vals[RANKING] = float(rel["total_points"]) if rel is not None else None
        ns[RANKING] = 1 if rel is not None else 0
        src = win.iloc[0] if len(win) else None
        duration = int(src["nominal_minutes"]) if src is not None else None
        vals[DURATION] = duration
        wrow = {"excel_row": r, "match_number": num, "round": ws.cell(r, header["Round"]).value,
                "team": team_wb, "row_kickoff_utc": m["Date"], "source_type": wtype, "n_matches": len(win),
                "source_match": describe(src) if src is not None else None,
                "source_kickoff_utc": src["kickoff"].isoformat() if src is not None else None,
                "source_fifa_match_id": src["fifa_match_id"] if src is not None else None,
                "source_extra_time": bool(src["extra_time"]) if src is not None else None,
                DURATION: duration,
                "count_scale_to_per90": 90 / src["nominal_minutes"] if src is not None else None,
                "saving_note": why.get("saving") if vals["saving"] is not None else None,
                "all_before_kickoff": bool((win["kickoff"] < kickoff).all()) if len(win) else True,
                "ranking_date_id": rel["date_id"] if rel is not None else None,
                "ranking_release_utc": rel["release_utc"] if rel is not None else None,
                "ranking_rank": int(rel["rank"]) if rel is not None else None}
        for col, (key, _) in COLUMNS.items():
            wrow[f"has_{key}"] = ns[key]
        win_rows.append(wrow)
        if src is not None:
            used_in.setdefault((src["team"], src["fifa_match_id"]), []).append(r)

        for col, (key, unit) in COLUMNS.items():
            if vals[key] is None or (isinstance(vals[key], float) and math.isnan(vals[key])):
                if key == RANKING:
                    reason = "team not found in any FIFA ranking release before kickoff"
                elif win.empty:
                    reason = "no previous match found"
                elif key == "free_kick" and not fk_ok:
                    reason = "free kicks won only reported by FIFA for WC matches; derived measure rejected"
                elif why.get(key) and why[key].startswith("undefined"):
                    reason = why[key]
                else:
                    reason = ("no source publishes this stat for the source match (FotMob, ESPN box, FIFA live "
                              "API): " + lacking(win, key) + ". Not replaced with an older match")
                miss_rows.append({"excel_row": r, "match_number": num, "team": team_wb, "column": col,
                                  "status": "missing", "source_match": wrow["source_match"], "reason": reason})
        out_rows.append((r, vals, ns, act, date_fixed))

    # ---- write sheet "in" (layout unchanged; values + header notes)
    for r, vals, ns, act, date_fixed in out_rows:
        c = ws.cell(r, header["Date"])
        c.value = dt.datetime(date_fixed.year, date_fixed.month, date_fixed.day)
        c.number_format = "dd/mm/yyyy"
        ws.cell(r, header["act_goals"]).value = act
        ws.cell(r, header["pred_goals"]).value = None
        if vals is None:
            continue
        for col, (key, _) in COLUMNS.items():
            x = vals[key]
            ws.cell(r, header[col]).value = None if x is None or (isinstance(x, float) and math.isnan(x)) \
                else round(float(x), 4)
        ws.cell(r, header[DURATION]).value = vals[DURATION]
    for col, (key, unit) in COLUMNS.items():
        ws.cell(1, header[col]).comment = Comment(
            f"{unit}. Value from ONE match: the team's previous completed match before kickoff (previous WC "
            f"match; for openers the team's last international friendly before the tournament, or its latest "
            f"friendly with published stats if the last one has none). No averaging. "
            f"{'Original header: ' + repr(RENAMED[col]) + '. ' if col in RENAMED else ''}"
            f"See 'definitions', 'row_windows'.",
            "data prep")
    ws.cell(1, header[DURATION]).comment = Comment(
        "Length of the source match (the one match this row's predictors come from), in nominal minutes: 90, "
        "or 120 if it went to extra time. A penalty shoot-out adds no minutes. Counts in this row are already "
        "per 90. New column (not in the original workbook). See 'row_windows'.", "data prep")
    ws.cell(1, header["saving"]).comment = Comment(
        "% of opponent shots on target saved (GK saves / opponent shots on target), from ONE match: the previous "
        "WC match, or for openers the team's last friendly with published stats. When the opponent had 0 shots "
        "on target the value is 0 (team rule), so 0 also covers 'faced no shot on target'; those rows are "
        "marked in row_windows.saving_note. Capped at 100 when a source reports more saves than shots on "
        "target (also marked). See 'definitions'.", "data prep")
    ws.cell(1, header["FIFA ranking points"]).comment = Comment(
        "FIFA men's world ranking total points from the last official release before kickoff (11 Jun 2026, "
        "10:00 UTC). Not from the source match; same value in every row of a team because FIFA published no "
        "ranking during the tournament. Replaces 'km running' (original header). See 'definitions'.", "data prep")
    ws.cell(1, header["act_goals"]).comment = Comment(
        "Team's goals in this match incl. extra time, excl. penalty shoot-out (FIFA).", "data prep")
    ws.cell(1, header["pred_goals"]).comment = Comment("Left blank on purpose (no model yet).", "data prep")
    ws.cell(1, header["Date"]).comment = Comment(
        "Kickoff date (UTC) as true date, dd/mm/yyyy. Originals logged in 'match_check'.", "data prep")

    # ---- supporting sheets
    match_check = build_match_check(cal, date_log)
    defs = definitions(fk_ok, fk_mae, fk_bias, fk_r, fk_exact, len(v))
    fk_sheet = v[["team", "opponent", "kickoff_utc", "wc_match_number", "fk_won_reported", "fk_derived",
                  "opp_fouls_committed", "opp_offsides"]].assign(diff=v["fk_derived"] - v["fk_won_reported"])
    fk_summary = pd.DataFrame([{"team": "SUMMARY", "opponent": f"n={len(v)}",
                                "kickoff_utc": f"MAE={fk_mae:.3f} bias={fk_bias:+.3f} r={fk_r:.3f} "
                                               f"exact={fk_exact:.1%}",
                                "wc_match_number": None,
                                "fk_won_reported": "accepted" if fk_ok else "rejected"}])
    raw_out = raw.drop(columns=["kickoff"])
    raw_out.insert(0, "used_in_rows", [", ".join(str(x) for x in used_in.get((t, mid), [])) or "not used"
                                       for t, mid in zip(raw["team"], raw["fifa_match_id"])])
    sheets = {
        "raw_history": raw_out,
        "row_windows": pd.DataFrame(win_rows),
        "definitions": defs,
        "sources": sources_table(),
        "missing_notes": pd.DataFrame(miss_rows, columns=["excel_row", "match_number", "team", "column", "status",
                                                          "source_match", "reason"]),
        "match_check": match_check,
        "freekick_validation": pd.concat([fk_summary, fk_sheet], ignore_index=True),
        "crosscheck": cross,
        "friendly_lists": lists,
        "fifa_ranking": ranking.drop(columns=["released"]),
    }
    for name, df in sheets.items():
        if name in wbook.sheetnames:
            del wbook[name]
        sh = wbook.create_sheet(name)
        sh.append(list(df.columns))
        for rec in df.itertuples(index=False):
            sh.append([None if (isinstance(x, float) and math.isnan(x)) else
                       (x.item() if hasattr(x, "item") else x) for x in rec])
        for i, colname in enumerate(df.columns, 1):
            width = min(47, max(12, len(str(colname)) + 6,
                                *(len(str(x)) + 2 for x in df.iloc[:200, i - 1].tolist())))
            sh.column_dimensions[get_column_letter(i)].width = width
        style_sheet(sh, "A2")
    style_sheet(ws, "A5")
    for col, width in IN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    wbook.save(WORKBOOK_OUT)
    print(f"wrote {WORKBOOK_OUT}")

    verify(WORKBOOK_OUT, cal, pd.DataFrame(win_rows), ranking)
    write_completeness(pd.DataFrame(win_rows), pd.DataFrame(miss_rows), out_rows, fk_ok, fk_mae, fk_r, fk_exact,
                       len(v), match_check, cross)


# Workbook look chosen by the user in Excel (16 Oct 2026), re-applied on every rebuild.
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F3864")
BODY_FONT = Font(name="Calibri", size=10)
ROW_LINE = Side(style="thin", color="FFD9D9D9")
IN_WIDTHS = {"A": 12.0, "B": 12.83, "C": 18.5, "D": 20.0, "E": 24.33, "F": 19.33, "G": 17.83, "H": 15.66, "I": 12.83,
             "J": 25.5, "K": 26.83, "L": 18.5, "M": 21.5, "N": 24.66, "O": 17.0, "P": 16.83, "Q": 16.0, "R": 14.83,
             "S": 26.5}


def style_sheet(sh, freeze):
    """Dark-blue bold header (22 pt row, filter), Calibri 10 body with light grey row lines, frozen panes."""
    for c in sh[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    sh.row_dimensions[1].height = 22
    sh.auto_filter.ref = f"A1:{get_column_letter(sh.max_column)}1"
    sh.freeze_panes = freeze
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="center")
            c.border = Border(top=ROW_LINE if c.row > 2 else Side(), bottom=ROW_LINE)


def build_match_check(cal, date_log):
    dl = pd.DataFrame(date_log)
    rows = []
    for num in sorted(cal):
        m = cal[num]
        sub = dl[dl["match_number"] == num]
        teams = sorted(sub["team"])
        fifa = sorted([m["Home"]["TeamName"][0]["Description"], m["Away"]["TeamName"][0]["Description"]])
        teams_fifa = sorted(WB_TO_FIFA.get(t, t) for t in teams)
        rt = m.get("ResultType")
        note = {1: "normal time", 2: "extra time + penalty shoot-out (shoot-out excluded from score)",
                3: "after extra time (score includes extra-time goals)"}.get(rt, f"ResultType {rt}")
        rows.append({"match_number": num, "workbook_round": "; ".join(sorted(set(sub["round"].astype(str)))),
                     "fifa_stage": m["StageName"][0]["Description"],
                     "workbook_teams": " v ".join(teams), "fifa_home": m["Home"]["TeamName"][0]["Description"],
                     "fifa_away": m["Away"]["TeamName"][0]["Description"],
                     "teams_match": teams_fifa == fifa,
                     "workbook_date_original": "; ".join(sorted(set(sub["original_value"]))),
                     "workbook_date_normalised": sub["normalised_date"].iloc[0],
                     "fifa_kickoff_utc": m["Date"], "fifa_local_kickoff": m["LocalDate"],
                     "date_matches_utc": bool(sub["matches_fifa_utc_date"].all()),
                     "score": f'{m["HomeTeamScore"]}-{m["AwayTeamScore"]}',
                     "penalties": (f'{m["HomeTeamPenaltyScore"]}-{m["AwayTeamPenaltyScore"]}'
                                   if m.get("HomeTeamPenaltyScore") is not None else None),
                     "result_type": note,
                     "date_fix": "; ".join(sorted(set(sub["how"]))),
                     "fifa_url": (f"https://www.fifa.com/en/match-centre/match/17/285023/{m['IdStage']}/"
                                  f"{m['IdMatch']}")})
    return pd.DataFrame(rows)


def definitions(fk_ok, mae, bias, r, exact, n):
    fk_text = (f"Derived for every match as opponent fouls committed + opponent offsides (each foul or offside "
               f"gives a free kick). Validated against FIFA's reported 'FreeKicks' on {n} WC team-matches: "
               f"MAE {mae:.2f}, bias {bias:+.2f}, r {r:.3f}, exact {exact:.0%}. "
               f"{'Accepted' if fk_ok else 'Rejected'} (rule: MAE <= {FK_MAX_MAE}, r >= {FK_MIN_R}). "
               "One definition is used for WC matches and friendlies so every row uses the same measure. "
               "FIFA's reported value is kept in raw_history.fk_won_reported. Not free-kick shots.")
    rows = [
        ("free_kicks_won_per90", "free kicks won", "per 90", "(opp fouls + opp offsides) * 90 / nominal min",
         "FIFA (WC), FotMob/ESPN (friendlies)", fk_text),
        ("Shot accurrancy", "shots on target / total shots", "%", "100 * on target / total shots",
         "FIFA (WC), FotMob/ESPN (friendlies)",
         "Accuracy, not conversion (goals/shots). Total shots include blocked shots. FIFA, FotMob and ESPN "
         "values were identical on every WC match checked."),
        ("Pass acurracy", "completed passes / attempted passes", "%", "100 * completed / attempted",
         "FotMob (ESPN fallback) for all matches",
         "FIFA's tracking count is ~3-6% higher than Opta-style counts (FotMob = ESPN), so FIFA passes are not mixed "
         "in; kept as raw_history.fifa_passes_att/cmp for reference."),
        ("possession", "share of ball possession", "%", "possession % in the source match",
         "FIFA live API (friendlies; ESPN/FotMob when FIFA has none), ESPN then FotMob (WC)",
         "Event-based possession (both teams sum to 100). FIFA WC tracking possession has an 'in contest' share "
         "and differs by up to ~4 points; kept as raw_history.fifa_possession_renorm (own/(own+opp)) for "
         "reference."),
        ("saving", "save percentage", "%", "100 * GK saves / opponent shots on target; 0 if the opponent had 0 "
         "shots on target; capped at 100", "FIFA (WC: GoalkeeperSavesOnTarget), FotMob/ESPN (friendlies)",
         "Save %, not save count. FIFA's 'GoalkeeperSaves' also counts non-on-target actions, so "
         "'GoalkeeperSavesOnTarget' is used (equals FotMob/ESPN saves). When the opponent had 0 shots on target "
         "(0 saves / 0 shots, mathematically undefined) the value is 0 by the team's rule (16 Oct 2026), so 0 "
         "means either 'saved none of the shots on target faced' or 'faced no shot on target'; the second case "
         "is marked in row_windows.saving_note. A value above 100 (source reports more saves than opponent "
         "shots on target; FotMob friendly data only) is capped at 100 and also marked (data check, 17 Oct 2026)."),
        ("fouls_committed_per90", "fouls committed", "per 90", "fouls committed * 90 / nominal min",
         "FIFA (WC: 'FoulsAgainst'), FotMob/ESPN (friendlies)",
         "FIFA's 'FoulsAgainst' is fouls committed by the team (confirmed: team FreeKicks = opponent FoulsAgainst "
         "+ opponent Offsides; equals FotMob/ESPN fouls committed)."),
        ("tackles_attempted_per90", "tackles attempted", "per 90", "tackles * 90 / nominal min",
         "FotMob (ESPN fallback) for all matches",
         "FIFA publishes no tackle count. FotMob 'Tackles' equals ESPN 'totalTackles' (all tackle attempts; "
         "ESPN 'effectiveTackles' = won tackles, kept in crosscheck)."),
        ("corners_per90", "corners won", "per 90", "corners * 90 / nominal min",
         "FIFA (WC), FotMob/ESPN (friendlies)", ""),
        ("FIFA ranking points", "FIFA/Coca-Cola men's world ranking total points", "points (2 dp)",
         "TotalPoints of the team in the latest release published before kickoff", "FIFA ranking API",
         "Replaces 'km running' (user request, 16 Oct 2026; same column). "
         "Release 11 Jun 2026 10:00 UTC (dateId id15136) precedes every kickoff (opener 11 Jun 19:00 UTC); the "
         "next release (20 Jul 2026) came after the final, so every row of a team has the same value. Not taken "
         "from the source match and not scaled per 90. Cross-check: the 20 Jul release's previous points equal "
         "the 11 Jun total points for all 48 teams. Distance covered stays in raw_history.distance_km."),
        ("goals_conceded_per90", "goals conceded in the source match", "per 90", "goals against * 90 / nominal min",
         "FIFA scores", "Extra-time goals included, shoot-out goals excluded, own goals count. Never the current "
         "match."),
        ("cards_per90", "yellow + red cards", "per 90", "(yellow + red) * 90 / nominal min",
         "FIFA bookings (players only); FotMob fallback",
         "Second-yellow dismissal = 1 yellow + 1 red (the second yellow is not counted twice). Yellows and reds "
         "kept separately in raw_history."),
        ("shots_per90", "total shots (on + off target + blocked)", "per 90", "shots * 90 / nominal min",
         "FIFA (WC), FotMob/ESPN (friendlies)", "Total attempts, not shots on target."),
        ("act_goals", "team goals in the current match", "goals", "FIFA score", "FIFA calendar API",
         "Includes extra-time goals, excludes penalty shoot-out goals."),
        ("pred_goals", "left blank", "", "", "", "No model trained yet."),
        (DURATION, "length of the source match", "minutes", "90, or 120 if extra time was played",
         "FIFA ResultType",
         "New column (user request, 16 Oct 2026). Shows whether the one match this row's predictors come from "
         "went to extra time. A penalty shoot-out adds no minutes. Per-90 columns are already scaled with it."),
        ("(all)", "nominal minutes", "minutes", "90, or 120 if extra time was played", "FIFA ResultType",
         "Used for every per-90 rate: counts from a 120-minute match are multiplied by 90/120; counts from a "
         "90-minute match are unchanged. Stoppage time ignored because friendly sources do not publish actual "
         "minutes."),
        ("(all)", "source match", "", "exactly ONE previous completed match before kickoff; no averaging", "",
         "Later rows: the team's previous WC 2026 match. Opener: the team's last international friendly before "
         "the opener (user decision 16 Oct 2026: one match type for every opener; qualifiers are not used). If "
         "that friendly has no published team statistics on any source, the opener uses the team's most recent "
         "friendly that has them (user decision 16 Oct 2026). Friendly stats: FotMob, ESPN fallback and "
         "cross-check, FIFA live API cards and possession. A stat missing from the chosen match stays blank; an "
         "older match is never used instead. Exception: FIFA ranking points (pre-match ranking)."),
    ]
    df = pd.DataFrame(rows, columns=["workbook_column", "measure", "unit", "formula", "source", "notes"])
    df.insert(1, "original_header", [RENAMED.get(c, "(new column)" if c == DURATION else c if c != "(all)" else "")
                                     for c in df["workbook_column"]])
    return df


def sources_table():
    acc = dt.date.today().isoformat()
    rows = [
        ("FIFA calendar API", "allowed (fifa.com)", "https://api.fifa.com/api/v3/calendar/matches?idCompetition=17"
         "&idSeason=285023", "fixtures, UTC kickoff, scores, extra time / penalties, act_goals", acc),
        ("FIFA team form API", "allowed (fifa.com)", "https://api.fifa.com/api/v3/teamform/{IdTeam}",
         "each team's friendly list (and last match of any competition) with dates and scores", acc),
        ("FIFA live match API", "allowed (fifa.com)", "https://api.fifa.com/api/v3/live/football/{IdMatch}",
         "bookings (cards) for all matches, friendly possession", acc),
        ("FIFA Data Hub team stats", "allowed (fifa.com match centre)",
         "https://fdh-api.fifa.com/v1/stats/match/{IdIFES}/teams.json",
         "WC: shots, on target, saves, fouls, corners, offsides, free kicks, distance, time played", acc),
        ("FIFA world ranking API", "allowed (fifa.com)", "https://inside.fifa.com/api/ranking-overview?locale=en"
         "&dateId={id15065|id15136|id15175}",
         f"FIFA ranking points (11 Jun 2026 release used); page {FIFA_RANKING_PAGE}", acc),
        ("FotMob match details", "OUTSIDE source (disclosed)", "https://www.fotmob.com/api/data/matchDetails?matchId=",
         "friendlies: shots, on target, saves, fouls, tackles, corners, offsides; all matches: passes, tackles; "
         "WC possession fallback", acc),
        ("FotMob team fixtures", "OUTSIDE source (disclosed)", "https://www.fotmob.com/api/data/teams?id=",
         "newest matches before each opener, to cross-check FIFA's friendly lists (friendly_lists sheet)", acc),
        ("ESPN summary API", "OUTSIDE source (disclosed)",
         "https://site.web.api.espn.com/apis/site/v2/sports/soccer/all/summary?event=",
         "cross-check; fallback when FotMob lacks a stat; WC possession (one decimal)", acc),
        ("Sofascore", "OUTSIDE source, checked; NOT used in this version",
         "https://www.sofascore.com/api/v1/event/12183390/statistics",
         "approved by the user for one qualifier (Cabo Verde v Eswatini) in the qualifier version (branch data-2); "
         "the friendly-only version uses no qualifier, so no Sofascore value is used. The saved response stays in "
         "raw_cache/sofascore/", "2026-10-16"),
        ("Flashscore", "OUTSIDE source, checked and rejected by user",
         "https://www.flashscore.com/",
         "not used: on 4 African qualifiers shots, shots on target and possession differed strongly from FotMob "
         "(e.g. Morocco v Congo shots 8 vs 18, possession 63 vs 77)", "2026-10-16"),
        ("FBref", "allowed, NOT usable", "https://fbref.com/en/",
         "Cloudflare human-verification blocks automated access; Opta-based advanced team stats removed from "
         "FBref on 20 Jan 2026", acc),
        ("The Stats Don't Lie", "allowed, NOT usable for per-match values",
         "https://www.thestatsdontlie.com/football/world-cup-2026/",
         "only tournament-level aggregates; no per-match team pages and no friendly coverage", acc),
    ]
    return pd.DataFrame(rows, columns=["source", "status", "endpoint", "used_for", "accessed"])


def verify(path, cal, windows, ranking):
    wb = openpyxl.load_workbook(path)
    ws = wb["in"]
    hdr = {c.value: c.column for c in ws[1]}
    assert ws.max_row == 209, ws.max_row
    assert ws.max_column == 19 and ws.cell(1, 19).value == DURATION, (ws.max_column, ws.cell(1, 19).value)
    assert all(new in hdr for new in RENAMED) and not any(old in hdr for old in RENAMED.values())
    nums = [ws.cell(r, hdr["MatchNumber"]).value for r in range(2, 210)]
    assert sorted(set(nums)) == list(range(1, 105)) and all(nums.count(n) == 2 for n in set(nums))
    for r in range(2, 210, 2):
        n = ws.cell(r, hdr["MatchNumber"]).value
        m = cal[n]
        goals = sorted([ws.cell(r, hdr["act_goals"]).value, ws.cell(r + 1, hdr["act_goals"]).value])
        assert goals == sorted([m["HomeTeamScore"], m["AwayTeamScore"]]), (n, goals)
        assert ws.cell(r, hdr["pred_goals"]).value is None
    if len(windows):
        assert windows["all_before_kickoff"].all()
        assert windows["n_matches"].isin([0, 1]).all()
        assert (pd.to_datetime(windows["ranking_release_utc"], utc=True)
                < pd.to_datetime(windows["row_kickoff_utc"], utc=True)).all()
        # openers (each team's first row) come from a friendly, every later row from the previous WC match
        opener = windows.index.isin(windows.sort_values("row_kickoff_utc").groupby("team").head(1).index)
        src = windows["source_match"].fillna("")
        assert windows.loc[opener, "source_type"].isin([WT_FR, WT_FR_EARLIER]).all()
        assert src[opener].str.contains(" FRI ").all()
        assert (windows.loc[~opener, "source_type"] == WT_WC).all() and src[~opener].str.contains(" WC#").all()
        # duration column: 90 or 120, 120 exactly when the source match went to extra time
        assert windows[DURATION].isin([90, 120]).all()
        assert ((windows[DURATION] == 120) == windows["source_extra_time"].astype(bool)).all()
        for row, d, has_sv in zip(windows["excel_row"], windows[DURATION], windows["has_saving"]):
            assert ws.cell(row, hdr[DURATION]).value == d, (row, d)
            # saving is never blank when the source match reports saves and opponent shots on target
            assert not has_sv or ws.cell(row, hdr["saving"]).value is not None, row
    june = ranking[ranking["date_id"] == "id15136"]
    pts = dict(zip(june["team"], june["total_points"]))
    for r in range(2, 210):
        sv = ws.cell(r, hdr["saving"]).value
        assert sv is None or 0 <= sv <= 100, (r, sv)
        team = ws.cell(r, hdr["Team"]).value
        x = ws.cell(r, hdr["FIFA ranking points"]).value
        if len(windows) == 208:
            assert x is not None and abs(x - pts[WB_TO_FIFA.get(team, team)]) < 1e-9, (r, team, x)
    print("verification passed")


def write_completeness(win, miss, out_rows, fk_ok, mae, r, exact, n_fk, match_check, cross):
    total = len(out_rows)
    filled = {col: 0 for col in COLUMNS}
    from_wc = {col: 0 for col in COLUMNS}
    stype = dict(zip(win["excel_row"], win["source_type"])) if len(win) else {}
    for r_, vals, ns, _, _ in out_rows:
        if vals is None:
            continue
        for col, (key, _) in COLUMNS.items():
            x = vals[key]
            if x is not None and not (isinstance(x, float) and math.isnan(x)):
                filled[col] += 1
                if stype.get(r_) == WT_WC:
                    from_wc[col] += 1
    lines = ["# Completeness summary", "",
             f"Rows: {total} team-match rows (104 matches). `act_goals` filled for all rows from FIFA. "
             "`pred_goals` left blank.", "",
             "Every predictor comes from exactly ONE previous match of the team: the previous WC match, or for "
             "a team's opener its last international friendly before the tournament (its latest friendly with "
             "published statistics if the last one has none). No averaging. Qualifiers are not used. "
             "Exception: `FIFA ranking points` (replaces "
             "`km running`) is the team's total points in the last official FIFA ranking before kickoff "
             "(11 Jun 2026 release, 10:00 UTC, before every match).", "",
             "| column | rows with value | from previous WC match | from friendly (opener) | rows blank |",
             "|---|---|---|---|---|"]
    for col, (key, _) in COLUMNS.items():
        if key == RANKING:
            lines.append(f"| {col} | {filled[col]} | n/a: 11 Jun 2026 FIFA ranking | n/a | "
                         f"{total - filled[col]} |")
            continue
        lines.append(f"| {col} | {filled[col]} | {from_wc[col]} | {filled[col] - from_wc[col]} | "
                     f"{total - filled[col]} |")
    dur = win[DURATION].value_counts().to_dict() if len(win) else {}
    lines.append(f"| {DURATION} | {sum(dur.values())} | 90: {dur.get(90, 0)}, 120: {dur.get(120, 0)} | "
                 f"(all openers' friendlies ended in normal time: {int((win.loc[win['source_type'] != WT_WC, DURATION] == 90).sum()) if len(win) else 0}"
                 f"/{int((win['source_type'] != WT_WC).sum()) if len(win) else 0}) | "
                 f"{total - sum(dur.values())} |")
    lines += ["",
              f"Free-kick derivation check: n={n_fk} WC team-matches, MAE {mae:.2f}, r {r:.3f}, exact {exact:.0%}: "
              f"{'accepted' if fk_ok else 'rejected'}.",
              f"Match identity: {int(match_check['teams_match'].sum())}/104 team pairs and "
              f"{int(match_check['date_matches_utc'].sum())}/104 dates agree with FIFA (dates read as day/month, "
              "UTC).", ""]
    raw = pd.read_csv(RAW_CSV)
    lines += ["## Which source supplied each raw value", "",
              "Counts of raw team-match values by provider (outside sources: FotMob, ESPN).", "",
              "| raw field | match type | source counts |", "|---|---|---|"]
    for f in ["shots_total", "shots_on_target", "passes_att", "possession_pct", "saves", "fouls_committed",
              "opp_fouls_committed", "tackles", "corners", "yellow_cards"]:
        for kind, g in raw.groupby("match_type"):
            vc = g[f"src_{f}"].fillna("missing").value_counts().to_dict()
            lines.append(f"| {f} | {kind} | " + ", ".join(f"{k}: {v}" for k, v in vc.items()) + " |")

    # data-driven facts for the decision notes
    earlier = win[win["source_type"] == WT_FR_EARLIER] if len(win) else win
    note = win["saving_note"].fillna("") if len(win) else pd.Series(dtype=str)
    zero, capped = note.str.startswith("0/0"), note.str.startswith("capped")
    n_zero = int(zero.sum())
    n_zero_open = int((zero & (win["source_type"] != WT_WC)).sum()) if len(win) else 0
    rejected = raw[raw["src_possession_pct"].fillna("").str.contains("rejected")]
    lists = pd.read_csv(INTERMEDIATE / "friendly_lists.csv")
    fifa1 = lists[(lists["source"] == "FIFA teamform (friendlies)") & (lists["rank"] == 1)].set_index("team")
    any1 = lists[lists["source"] == "FIFA teamform (any competition)"].set_index("team")
    fm1 = lists[(lists["source"] == "FotMob team fixtures") & (lists["rank"] == 1)].set_index("team")
    last_is_friendly = int((any1["fifa_match_id"] == fifa1["fifa_match_id"].reindex(any1.index)).sum())
    same_day = [t for t in fifa1.index if t in fm1.index and abs(
        (pd.Timestamp(fm1.loc[t, "utc"], tz="UTC") - pd.Timestamp(fifa1.loc[t, "utc"], tz="UTC")).total_seconds())
        < 36 * 3600]
    fm_newer = [t for t in fifa1.index if t in fm1.index and t not in same_day
                and pd.Timestamp(fm1.loc[t, "utc"], tz="UTC") > pd.Timestamp(fifa1.loc[t, "utc"], tz="UTC")]
    fm_other = sorted(set(fifa1.index) - set(same_day) - set(fm_newer))
    lines += ["", "## Decisions made while building", "",
              "- Openers use friendlies only (user decision after a team meeting, 16 Oct 2026): mixing qualifiers "
              "and friendlies made opener rows come from different match types. Each team's first row now uses "
              "its last international friendly before the tournament; the later rows keep the previous WC match. "
              f"For {last_is_friendly} of {len(any1)} teams the last match of any competition before the opener "
              "(FIFA team form) is that friendly.",
              "- If the last friendly has no published team statistics on FIFA, FotMob or ESPN, the opener uses the "
              "team's most recent friendly that has them (user decision, 16 Oct 2026). Applied to "
              f"{len(earlier)} team(s): "
              + ("; ".join(f"{t} ({m})" for t, m in zip(earlier["team"], earlier["source_match"])) or "none") + ".",
              f"- `saving` is 0 when the opponent had 0 shots on target (team rule, 16 Oct 2026). Before this rule "
              "every blank saving cell was such a 0/0 case (15 of 15 checked); none was missing data. "
              f"This build codes {n_zero} row(s) this way ({n_zero_open} opener(s)); they are marked in "
              "row_windows.saving_note. A 0 therefore means either 'saved none of the shots on target faced' or "
              "'faced no shot on target'.",
              f"- `saving` is capped at 100 when the source reports more saves than opponent shots on target "
              f"(data check, 17 Oct 2026). Applied to {int(capped.sum())} row(s)"
              + (": " + "; ".join(f"row {r_} {t} ({m})" for r_, t, m in
                                  zip(win.loc[capped, "excel_row"], win.loc[capped, "team"],
                                      win.loc[capped, "source_match"])) if capped.any() else "")
              + "; marked in row_windows.saving_note.",
              "- Headers renamed in this copy only (user request, 16 Oct 2026): "
              + ", ".join(f"`{old}` -> `{new}`" for new, old in RENAMED.items())
              + f". New last column `{DURATION}`: the source match's length, 90 or 120 (extra time).",
              f"- FotMob team fixtures as a check on FIFA's friendly lists: for {len(same_day)} of {len(fifa1)} teams "
              "FotMob's newest match before the opener is FIFA's last friendly. "
              + (f"FotMob has a newer match for: {', '.join(fm_newer)}. " if fm_newer else
                 "FotMob never has a newer match, so FIFA's list omits no later friendly. ")
              + (f"For {', '.join(fm_other)} FotMob does not list FIFA's last friendly at all (its newest match is "
                 "an older one); those are the matches without statistics." if fm_other else ""),
              "- `km running` was replaced by `FIFA ranking points` (user request, 16 Oct 2026). The user chose "
              "the official 11 Jun 2026 release over recomputing points after each match (a value FIFA never "
              "published) or the release before the source match (mixed dates). FIFA published no ranking "
              "between 11 Jun and 20 Jul 2026, so a team has the same value in all its rows.",
              "- Passes and possession use one event-based definition for every match. FIFA's WC tracking data "
              "counts ~3-6% more passes and leaves an 'in contest' share of possession, so mixing it with "
              "friendly data would make opener rows and later rows measure different things. FIFA values are "
              "kept in raw_history.",
              "- Free kicks won are derived (opponent fouls + opponent offsides) for WC matches too, for the same "
              "reason. FIFA's reported value is kept and the validation is in sheet 'freekick_validation'.",
              f"- FIFA's live-API possession is rejected when it is >10 points from both ESPN and FotMob: "
              f"{len(rejected)} record(s) in this build"
              + (": " + "; ".join(f"{t} v {o} ({k[:10]})" for t, o, k in
                                  zip(rejected["team"], rejected["opponent"], rejected["kickoff_utc"])) if len(rejected)
                 else "") + ".",
              "- Saves use FIFA 'GoalkeeperSavesOnTarget'; FIFA 'GoalkeeperSaves' also counts other keeper "
              "actions and can exceed opponent shots on target.",
              ""]
    if len(miss):
        miss = miss.assign(reason=miss["reason"].str.replace(r":\s.*$", "", regex=True))
        top = miss.groupby(["column", "reason"]).size().reset_index(name="rows").sort_values("rows", ascending=False)
        lines += ["## Blank values by reason", "",
                  "Row-level detail (which matches lack which stat) is in sheet 'missing_notes'.", "",
                  "| column | reason | rows |", "|---|---|---|"]
        lines += [f"| {a} | {b} | {c} |" for a, b, c in top.itertuples(index=False)]
    else:
        lines += ["## Blank values", "", "None: every predictor cell is filled.", ""]
    (OUTPUT / "COMPLETENESS.md").write_text("\n".join(lines) + "\n")
    print("wrote COMPLETENESS.md")


if __name__ == "__main__":
    sys.exit(main())
